"""Test cho Exercise.workbook(): đọc công thức + giá trị đã tính lại bằng LibreOffice."""

import shutil
from pathlib import Path

import openpyxl
import pytest

import testkit


def make_exercise(tmp_path: Path, monkeypatch, solution: bool = False):
    root = tmp_path / "tracks" / "t" / "m" / "exercises" / "01-x"
    (root / "tests").mkdir(parents=True)
    starter = root / "starter"
    starter.mkdir()
    sol = tmp_path / "tracks" / "t" / "m" / "solutions" / "01-x"
    sol.mkdir(parents=True)
    monkeypatch.setattr(testkit, "ROOT", tmp_path)
    if solution:
        monkeypatch.setenv("SOLUTION", "1")
    else:
        monkeypatch.delenv("SOLUTION", raising=False)
    return testkit.Exercise(str(root / "tests" / "test_x.py")), starter, sol


def write_book(path: Path, total_cell="B1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for i, v in enumerate((1, 2, 3), start=1):
        ws.cell(row=i, column=1, value=v)
    ws[total_cell] = "=SUM(A1:A3)"
    ws["C1"] = 6
    wb.save(path)


def test_formula_and_value(tmp_path, monkeypatch):
    ex, starter, _ = make_exercise(tmp_path, monkeypatch)
    write_book(starter / "b.xlsx")
    wb = ex.workbook("b.xlsx")
    assert wb.sheetnames == ["Data"]
    assert wb.formula("Data", "B1") == "=SUM(A1:A3)"
    assert wb.formula("Data", "C1") is None  # gõ cứng
    assert wb.formula("Data", "A1") is None
    if not shutil.which("soffice"):
        pytest.skip("không có soffice")
    assert wb.value("Data", "B1") == 6
    assert wb.pivot_tables("Data") == []


def test_value_recomputes_stale_cache(tmp_path, monkeypatch):
    # openpyxl không ghi giá trị cache; value() phải tính lại thay vì trả None
    ex, starter, _ = make_exercise(tmp_path, monkeypatch)
    write_book(starter / "b.xlsx")
    if not shutil.which("soffice"):
        pytest.skip("không có soffice")
    assert ex.workbook("b.xlsx").value("Data", "B1") == 6


def test_solution_env_points_to_solutions(tmp_path, monkeypatch):
    ex, starter, sol = make_exercise(tmp_path, monkeypatch, solution=True)
    write_book(sol / "b.xlsx", total_cell="B2")
    wb = ex.workbook("b.xlsx")
    assert wb.formula("Data", "B2") == "=SUM(A1:A3)"
    assert wb.formula("Data", "B1") is None


def test_missing_soffice_skips(tmp_path, monkeypatch):
    ex, starter, _ = make_exercise(tmp_path, monkeypatch)
    write_book(starter / "b.xlsx")
    monkeypatch.setattr(shutil, "which", lambda name, *a, **k: None)
    with pytest.raises(pytest.skip.Exception, match="LibreOffice"):
        ex.workbook("b.xlsx").value("Data", "B1")
    # formula không cần soffice
    assert ex.workbook("b.xlsx").formula("Data", "B1") == "=SUM(A1:A3)"


def test_missing_workbook_fails_clearly(tmp_path, monkeypatch):
    ex, *_ = make_exercise(tmp_path, monkeypatch)
    with pytest.raises(pytest.fail.Exception, match="Không thấy"):
        ex.workbook("nope.xlsx")
