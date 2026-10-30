"""Chấm bike_pivot.xlsx: cột Year, pivot table (mnth × Year, Sum of cnt), bảng YoY bằng SUMIFS và một chart 2 series.

Pivot được nhận ra từ định nghĩa pivot trong tệp .xlsx (LibreOffice và Excel đều ghi lại khi lưu dạng Excel);
giá trị kiểm trong vùng kết quả của pivot."""

import re
from collections import defaultdict

import openpyxl
import pytest
from openpyxl.utils import range_boundaries

from testkit import Exercise

ex = Exercise(__file__)
BOOK = "bike_pivot.xlsx"
FIRST, LAST = 2, 732


def norm(formula):
    return re.sub(r"_XLFN\.|_XLWS\.|COM\.MICROSOFT\.", "", (formula or "").upper())


def need_formula(wb, sheet, cell, *families):
    f = wb.formula(sheet, cell)
    if f is None:
        pytest.fail(f"{sheet}!{cell} chưa có công thức: ô đang trống hoặc gõ cứng giá trị. Điền công thức bắt đầu bằng =")
    for options in families:
        if not any(name in norm(f) for name in options):
            pytest.fail(f"{sheet}!{cell} = {f}: bài này yêu cầu dùng {' hoặc '.join(options)}")


@pytest.fixture(scope="module")
def wb():
    return ex.workbook(BOOK)


@pytest.fixture(scope="module")
def book(wb):
    return openpyxl.load_workbook(wb.path)  # đọc định nghĩa pivot và chart


@pytest.fixture(scope="module")
def monthly(wb):
    out = defaultdict(int)
    for r in range(FIRST, LAST + 1):
        out[(wb.value("Days", f"A{r}").year, wb.value("Days", f"D{r}"))] += wb.value("Days", f"L{r}")
    return out


def test_cot_year(wb):
    for r in range(FIRST, LAST + 1):
        need_formula(wb, "Days", f"M{r}", ("YEAR",))
        assert wb.value("Days", f"M{r}") == wb.value("Days", f"A{r}").year, f"Days!M{r}"


def test_bang_yoy(wb, monthly):
    for m in range(1, 13):
        r = m + 1
        need_formula(wb, "YoY", f"B{r}", ("SUMIFS",))
        need_formula(wb, "YoY", f"C{r}", ("SUMIFS",))
        need_formula(wb, "YoY", f"D{r}", ("/",))
        assert wb.value("YoY", f"B{r}") == monthly[(2011, m)], f"YoY!B{r}"
        assert wb.value("YoY", f"C{r}") == monthly[(2012, m)], f"YoY!C{r}"
        assert wb.value("YoY", f"D{r}") == pytest.approx(monthly[(2012, m)] / monthly[(2011, m)] - 1), f"YoY!D{r}"
    for c in "BCD":
        need_formula(wb, "YoY", f"{c}14", ("SUM", "/"))
    assert (wb.value("YoY", "B14"), wb.value("YoY", "C14")) == (1243103, 2049576)
    assert wb.value("YoY", "D14") == pytest.approx(0.6488, abs=1e-4)
    assert wb.value("YoY", "C10") == 218573          # tháng 9/2012, tháng cao nhất


def find_pivots(book):
    return [(name, p) for name in book.sheetnames for p in book[name]._pivots]


def test_pivot_table(wb, book):
    pivots = find_pivots(book)
    assert pivots, ("Chưa thấy pivot table nào. Tạo bằng Insert > Pivot Table (LibreOffice) / Insert > PivotTable (Excel) "
                    "rồi lưu dạng .xlsx")
    problems = []
    for sheet, p in pivots:
        fields = [f.name for f in p.cache.cacheFields]
        rows = [fields[f.x] for f in p.rowFields if f.x >= 0]
        cols = [fields[f.x] for f in p.colFields if f.x >= 0]
        data = [(fields[d.fld], d.subtotal) for d in p.dataFields]
        if "mnth" not in rows:
            problems.append(f"{sheet}: hàng của pivot là {rows}, cần mnth")
        elif not ({"Year", "yr"} & set(cols)):
            problems.append(f"{sheet}: cột của pivot là {cols}, cần Year")
        elif ("cnt", "sum") not in data:
            problems.append(f"{sheet}: giá trị của pivot là {data}, cần Sum of cnt")
        else:
            min_col, min_row, max_col, max_row = range_boundaries(p.location.ref)
            values = {wb.value(sheet, f"{openpyxl.utils.get_column_letter(c)}{r}")
                      for r in range(max(1, min_row - 2), max_row + 1) for c in range(min_col, max_col + 1)}
            missing = {218573, 38189, 1243103, 2049576} - values
            if not missing:
                return
            problems.append(f"{sheet}: vùng kết quả pivot thiếu các giá trị {sorted(missing)}")
    pytest.fail("Pivot chưa đúng: " + "; ".join(problems))


def test_chart_hai_series(book):
    charts = [(name, c) for name in book.sheetnames for c in book[name]._charts]
    assert charts, "Chưa có chart nào trong workbook (Insert > Chart)"
    kinds = [(name, type(c).__name__, len(c.series)) for name, c in charts]
    assert any(k in ("BarChart", "LineChart", "AreaChart") and n >= 2 for _, k, n in kinds), (
        f"Cần một chart cột/đường có 2 series (2011 và 2012); đang có {kinds}")
