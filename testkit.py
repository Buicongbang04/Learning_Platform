"""Tiện ích dùng chung cho test của mọi bài tập.

Mỗi bài tập có `starter/` (bạn sửa) và `tests/`. Lời giải mẫu nằm ở thư mục `solutions/`
song song với `exercises/`, cùng tên tệp với `starter/`.

    uv run pytest <thư mục bài tập> -q              # kiểm tra code của bạn trong starter/
    SOLUTION=1 uv run pytest <thư mục bài tập> -q   # kiểm tra lời giải mẫu

Trong test:

    from testkit import Exercise
    ex = Exercise(__file__)
    transform = ex.module("transform.py")   # import starter/transform.py (hoặc lời giải)
    analysis = ex.notebook("analysis.ipynb") # chạy code cell (trừ cell tag `explore`), dùng như module
    sql = ex.read("query.sql")              # đọc tệp văn bản
    ex.require_service("postgres")          # skip nếu dịch vụ Docker chưa chạy
    wb = ex.workbook("report.xlsx")         # bài Excel: wb.formula("Sheet", "B2"), wb.value("Sheet", "B2")
"""

from __future__ import annotations

import importlib.util
import linecache
import os
import re
import shutil
import signal
import socket
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from types import ModuleType
from typing import Any

import pytest

ROOT = Path(__file__).resolve().parent


def _load_env_file(path: Path) -> None:
    """Đọc setup/.env (cổng đã đổi khi trùng dự án khác). Biến môi trường thật vẫn được ưu tiên."""
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


_load_env_file(ROOT / "setup" / ".env")

# Cổng mặc định khớp setup/docker-compose.yml. Ghi đè bằng biến môi trường nếu cần.
SERVICES = {
    "postgres": ("POSTGRES_HOST", "localhost", "POSTGRES_PORT", 5432),
    "mongo": ("MONGO_HOST", "localhost", "MONGO_PORT", 27017),
    "redis": ("REDIS_HOST", "localhost", "REDIS_PORT", 6379),
    "cassandra": ("CASSANDRA_HOST", "localhost", "CASSANDRA_PORT", 9042),
    "neo4j": ("NEO4J_HOST", "localhost", "NEO4J_PORT", 7687),
    "minio": ("MINIO_HOST", "localhost", "MINIO_PORT", 9000),
    "kafka": ("KAFKA_HOST", "localhost", "KAFKA_PORT", 9094),
    "metabase": ("METABASE_HOST", "localhost", "METABASE_PORT", 3000),
    "airflow": ("AIRFLOW_HOST", "localhost", "AIRFLOW_PORT", 8080),
    "localstack": ("LOCALSTACK_HOST", "localhost", "LOCALSTACK_PORT", 4566),
}

def _port(name: str) -> int:
    return int(os.environ.get(SERVICES[name][2], SERVICES[name][3]))


POSTGRES_DSN = os.environ.get("POSTGRES_DSN", f"postgresql://de:de@localhost:{_port('postgres')}/de")
MONGO_URI = os.environ.get("MONGO_URI", f"mongodb://localhost:{_port('mongo')}")
REDIS_URL = os.environ.get("REDIS_URL", f"redis://localhost:{_port('redis')}/0")
NEO4J_URI = os.environ.get("NEO4J_URI", f"bolt://localhost:{_port('neo4j')}")
NEO4J_AUTH = ("neo4j", os.environ.get("NEO4J_PASSWORD", "password123"))
MINIO_ENDPOINT = os.environ.get("MINIO_ENDPOINT", f"http://localhost:{_port('minio')}")
MINIO_AUTH = ("minioadmin", "minioadmin")
LOCALSTACK_ENDPOINT = os.environ.get("LOCALSTACK_ENDPOINT", f"http://localhost:{_port('localstack')}")
KAFKA_BOOTSTRAP = os.environ.get("KAFKA_BOOTSTRAP", "localhost:9094")


def use_solution() -> bool:
    return os.environ.get("SOLUTION") == "1"


def service_up(name: str, timeout: float = 0.5) -> bool:
    host_env, host_default, port_env, port_default = SERVICES[name]
    host = os.environ.get(host_env, host_default)
    port = int(os.environ.get(port_env, port_default))
    try:
        with socket.create_connection((host, port), timeout=timeout):
            return True
    except OSError:
        return False


# Dòng magic `%...`, shell `!...` hoặc gán từ shell `x = !ls`; nhóm 1 giữ thụt lề.
_MAGIC = re.compile(r"^(\s*)(?:[%!]|[\w, ]+=\s*[%!]).*$")


def _unique_name(path: Path) -> str:
    return "ex_" + "_".join(path.relative_to(ROOT).with_suffix("").parts).replace("-", "_")


def _load(path: Path, unique: str, execute, spec=None) -> ModuleType:
    if unique in sys.modules:
        return sys.modules[unique]
    module = importlib.util.module_from_spec(spec) if spec else ModuleType(unique)
    module.__file__ = str(path)
    sys.path.insert(0, str(path.parent))
    try:
        sys.modules[unique] = module
        execute(module)
    except BaseException:
        del sys.modules[unique]
        raise
    finally:
        sys.path.remove(str(path.parent))
    return module


_RECALC_CACHE: dict[tuple, Any] = {}  # (đường dẫn, mtime_ns, size) -> openpyxl workbook đã tính lại


def _open_workbook(path: Path, **kwargs):
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        return openpyxl.load_workbook(path, **kwargs)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        pytest.fail(f"Không mở được workbook {path.name} ({exc}) — lưu lại tệp/đóng Excel rồi chạy lại")


def _recalculate(path: Path):
    """Chép sang thư mục tạm, để LibreOffice tính lại rồi đọc giá trị (data_only)."""
    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("Cần LibreOffice (soffice) để tính lại công thức")
    with tempfile.TemporaryDirectory(prefix="wbcalc-") as tmp:
        src = Path(tmp) / "in" / path.name
        out = Path(tmp) / "out"
        src.parent.mkdir()
        out.mkdir()
        shutil.copy(path, src)
        cmd = [
            soffice, f"-env:UserInstallation=file://{tmp}/profile", "--headless",
            "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src),
        ]
        proc = subprocess.Popen(
            cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, start_new_session=True
        )
        try:
            _, stderr = proc.communicate(timeout=120)
        except subprocess.TimeoutExpired:
            os.killpg(proc.pid, signal.SIGKILL)  # soffice sinh process con (soffice.bin)
            _, stderr = proc.communicate()
            pytest.fail(f"LibreOffice quá 120 giây khi tính lại {path.name}. stderr: {stderr.strip()}")
        result = out / src.name
        if proc.returncode != 0 or not result.exists():
            pytest.fail(
                f"LibreOffice không tính lại được {path.name} (exit {proc.returncode}, "
                f"{'có' if result.exists() else 'không có'} tệp kết quả). stderr: {stderr.strip()}"
            )
        return _open_workbook(result, data_only=True)


class WorkbookView:
    """Workbook .xlsx của người học: đọc công thức (không cần LibreOffice) và giá trị đã tính lại."""

    def __init__(self, path: Path):
        self.path = path
        self._formulas = None

    def _formula_book(self):
        if self._formulas is None:
            self._formulas = _open_workbook(self.path)  # data_only=False: giữ chuỗi công thức
        return self._formulas

    @property
    def sheetnames(self) -> list[str]:
        return list(self._formula_book().sheetnames)

    def _sheet(self, book, sheet: str):
        if sheet not in book.sheetnames:
            pytest.fail(f"Workbook {self.path.name} không có sheet '{sheet}'. Có: {book.sheetnames}")
        return book[sheet]

    def formula(self, sheet: str, cell: str) -> str | None:
        """Chuỗi công thức (`"=SUM(A1:A3)"`) nếu ô là công thức, ngược lại `None` (ô gõ cứng hoặc trống)."""
        c = self._sheet(self._formula_book(), sheet)[cell]
        if c.data_type != "f":
            return None
        raw = c.value
        return raw if isinstance(raw, str) else getattr(raw, "text", None)  # ArrayFormula có .text

    def pivot_tables(self, sheet: str) -> list:
        return list(getattr(self._sheet(self._formula_book(), sheet), "_pivots", []))

    def value(self, sheet: str, cell: str) -> Any:
        """Giá trị đã tính lại bằng LibreOffice (cache cấp module theo đường dẫn + mtime + size)."""
        stat = self.path.stat()
        key = (str(self.path.resolve()), stat.st_mtime_ns, stat.st_size)
        if key not in _RECALC_CACHE:
            _RECALC_CACHE[key] = _recalculate(self.path)
        return self._sheet(_RECALC_CACHE[key], sheet)[cell].value


class Exercise:
    """Một bài tập, xác định từ đường dẫn tệp test đang chạy."""

    def __init__(self, test_file: str):
        tests_dir = Path(test_file).resolve().parent
        self.dir = tests_dir.parent
        self.fixtures = tests_dir / "fixtures"
        parts = list(self.dir.parts)
        if use_solution():
            index = len(parts) - 1 - parts[::-1].index("exercises")
            parts[index] = "solutions"
            self.target = Path(*parts)
        else:
            self.target = self.dir / "starter"

    def path(self, name: str) -> Path:
        path = self.target / name
        if not path.exists() and use_solution() and "advanced" in self.dir.parts:
            pytest.skip("Bài advanced chỉ có SOLUTION.md (hướng làm + rubric), không có code lời giải để chạy")
        if not path.exists():
            pytest.fail(f"Không thấy tệp {path.relative_to(ROOT)}")
        return path

    def workbook(self, name: str) -> WorkbookView:
        """Workbook .xlsx của bài (starter hoặc lời giải với SOLUTION=1)."""
        return WorkbookView(self.path(name))

    def read(self, name: str) -> str:
        return self.path(name).read_text(encoding="utf-8")

    def module(self, name: str) -> ModuleType:
        path = self.path(name)
        spec = importlib.util.spec_from_file_location(_unique_name(path), path)
        return _load(path, spec.name, lambda module: spec.loader.exec_module(module), spec)

    def notebook(self, name: str) -> ModuleType:
        """Chạy các code cell của notebook (trừ cell gắn tag `explore`) rồi trả về như một module.

        Dòng magic/shell (`%time ...`, `!ls`, `x = !ls`) thành `pass`, cell magic (`%%bash`) bị bỏ,
        `display()` in ra như `print`. Output đã lưu trong notebook không ảnh hưởng. Lỗi được báo theo
        cell: `analysis.ipynb[cell 5]`, dòng tính trong cell.
        """
        import nbformat

        path = self.path(name)
        book = nbformat.read(path, as_version=4)
        cells = []
        for number, cell in enumerate(book.cells, start=1):
            if cell.cell_type != "code" or "explore" in cell.get("metadata", {}).get("tags", []):
                continue
            if cell.source.lstrip().startswith("%%"):
                continue
            source = "\n".join(_MAGIC.sub(r"\1pass", line) for line in cell.source.splitlines()) + "\n"
            filename = f"{path}[cell {number}]"
            linecache.cache[filename] = (len(source), None, source.splitlines(keepends=True), filename)
            cells.append(compile(source, filename, "exec"))
        os.environ.setdefault("MPLBACKEND", "Agg")

        def execute(module: ModuleType) -> None:
            module.__dict__.setdefault("display", print)
            for code in cells:
                exec(code, module.__dict__)

        return _load(path, _unique_name(path), execute)

    @staticmethod
    def dataset(name: str) -> Path:
        """Thư mục bộ dữ liệu thật `data/raw/<name>/`; skip test nếu chưa tải."""
        path = ROOT / "data" / "raw" / name
        if not path.is_dir() or not any(path.iterdir()):
            pytest.skip(f"Chưa tải bộ {name}. Chạy: uv run python data/download_dataset.py {name}")
        return path

    @staticmethod
    def require_service(name: str) -> None:
        if not service_up(name):
            pytest.skip(
                f"Dịch vụ {name} chưa chạy. Bật bằng: docker compose -f setup/docker-compose.yml up -d {name}"
            )
